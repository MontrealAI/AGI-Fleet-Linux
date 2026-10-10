#!/usr/bin/env python3
import json,subprocess,time,os,threading,resource
from pathlib import Path

class CheckFailure(Exception):
 def __init__(self,stage,command,result):
  self.detail={'stage':stage,'command':command,'exitCode':result.returncode,'stdout':result.stdout[-16000:],'stderr':result.stderr[-16000:],'desktopLog':'/tmp/agi-desktop/desktop.log','appsLog':'/tmp/agi-desktop/apps.log'}
  super().__init__(stage)

def command(stage,args,timeout=30):
 try:r=subprocess.run(args,env={**os.environ,'DISPLAY':':99'},capture_output=True,text=True,timeout=timeout)
 except subprocess.TimeoutExpired as e:
  r=subprocess.CompletedProcess(args,124,(e.stdout or b'').decode(errors='replace') if isinstance(e.stdout,bytes) else e.stdout or '',(e.stderr or b'').decode(errors='replace') if isinstance(e.stderr,bytes) else e.stderr or '')
 if r.returncode:raise CheckFailure(stage,args,r)
 return r.stdout

def desktop(*args):
 return json.loads(command('desktop-'+args[0],['agi-desktop',*args]))

def geometry():
 command('x-server-ready',['xdpyinfo','-display',':99'])
 windows=command('fixture-window',['xdotool','search','--onlyvisible','--name','^AGI local input check$']).splitlines()
 for window in windows:
  raw=command('fixture-geometry',['xdotool','getwindowgeometry','--shell',window])
  v=dict(line.split('=',1) for line in raw.splitlines() if '=' in line)
  if int(v.get('WIDTH',0))==500 and int(v.get('HEIGHT',0))==220:return {k:int(v[k]) for k in ['X','Y','WIDTH','HEIGHT']}
 raise RuntimeError('SELFTEST_FIXTURE_GEOMETRY_NOT_READY')

def stable_frame(observe=desktop,locate=geometry,clock=time.monotonic,pause=time.sleep,timeout=15,consecutive=3):
 end=clock()+timeout;last=None;stable=0;last_error=None
 while clock()<end:
  try:
   loc=locate();frame=observe('screenshot');key=(tuple(sorted(loc.items())),frame['frame'])
   stable=stable+1 if key==last else 1;last=key
   if stable>=consecutive:return frame,loc
  except (CheckFailure,RuntimeError,KeyError,ValueError) as e:last=None;stable=0;last_error=e
  pause(.15)
 if isinstance(last_error,CheckFailure):raise last_error
 raise RuntimeError('SELFTEST_FIXTURE_DID_NOT_SETTLE')

def fixture_action(action,arguments,observe=desktop,acquire=stable_frame,attempts=3):
 for attempt in range(attempts):
  frame,loc=acquire();args=arguments(loc) if callable(arguments) else arguments
  try:return observe(action,'--frame',frame['frame'],*args)
  except CheckFailure as e:
   try:code=json.loads(e.detail['stdout']).get('error')
   except (ValueError,AttributeError):raise e
   if code!='DESKTOP_SCENE_CHANGED_OBSERVE_AGAIN' or attempt+1==attempts:raise
 raise RuntimeError('SELFTEST_ACTION_ATTEMPTS_EXHAUSTED')

def read_fixture():
    value=json.loads(Path('/workspace/selftest/fixture-state.json').read_text())
    if value.get('schema')!='agi-synthetic-fixture/v1' or not value.get('ready') or time.monotonic()-value['at']>1:
        raise RuntimeError('SELFTEST_FIXTURE_NOT_READY')
    return value

def ready_state(state,focus=False,typed=False):
    return bool(state.get('ready') and (not focus or state.get('focus')=='entry') and
                (not typed or state.get('textMatches') is True) and state.get('caret',{}).get('off')==0)

def quiescent_frame(observe=desktop,state=read_fixture,clock=time.monotonic,pause=time.sleep,quiet=1.2,timeout=20,focus=False,typed=False):
    deadline=clock()+timeout;last=None;since=None
    while clock()<deadline:
        try:
            s=state()
            if not ready_state(s,focus,typed):last=since=None;pause(.05);continue
            frame=observe('screenshot');now=clock();key=(frame['frame'],tuple(sorted(s['geometry'].items())),s.get('focus'),s.get('textMatches'),s.get('selected'))
            if key!=last:last=key;since=now
            if since is not None and now-since>=quiet and now-s['lastVisualEvent']>=quiet:
                current=state()
                if not ready_state(current,focus,typed) or current['lastVisualEvent']!=s['lastVisualEvent']:last=since=None;continue
                return frame,current
        except (FileNotFoundError,ValueError,KeyError,RuntimeError):last=since=None
        pause(.05)
    raise RuntimeError('SELFTEST_FIXTURE_DID_NOT_SETTLE')

def pixel_diff(before,after):
    from PIL import ImageChops
    a=before.convert('RGB');b=after.convert('RGB')
    if a.size!=b.size:raise RuntimeError('SELFTEST_IMAGE_GEOMETRY')
    d=ImageChops.difference(a,b);return {'changedPixels':sum(pixel!=(0,0,0) for pixel in d.getdata()),'bbox':d.getbbox()}

class FixtureActions:
    def __init__(self,capture=False):self.attempts=[];self.capture=capture;self.images=[]
    def perform(self,stage,action,target=None,arguments=None,focus=False,typed=False):
        from PIL import Image
        for number in range(1,4):
            frame,state=quiescent_frame(focus=focus,typed=typed)
            geom=state[target+'Geometry'] if target else state['geometry']
            args=['--x',str(geom['X']+geom['WIDTH']//2),'--y',str(geom['Y']+geom['HEIGHT']//2)] if target else arguments or []
            before=Image.open('/workspace/.agi-desktop/screen.png').copy()
            record={'stage':stage,'action':action,'attempt':number,'frameBefore':frame['frame'],'observedAt':time.monotonic(),'geometry':geom,
                    'fixtureWindow':state['windowID'],'focusedWidget':state['entryID'] if state['focus']=='entry' else None,'focus':state['focus'],
                    'textMatches':state['textMatches'],'selected':state['selected'],'caret':state['caret'],'events':state['events']}
            self.attempts.append(record)
            record['dispatchStartedAt']=time.monotonic()
            try:
                result=desktop(action,'--frame',frame['frame'],*args);record.update(outcome='completed',exitCode=0,finishedAt=time.monotonic(),frameAfter=result.get('frame'));return result
            except CheckFailure as e:
                try:code=json.loads(e.detail['stdout']).get('error')
                except (ValueError,AttributeError):code='UNSTRUCTURED_ACTION_FAILURE'
                known=code=='DESKTOP_SCENE_CHANGED_OBSERVE_AGAIN' and e.detail['exitCode']==1
                record.update(outcome='rejected-before-dispatch' if known else 'ambiguous-failure',error=code,exitCode=e.detail['exitCode'],finishedAt=time.monotonic())
                if known:
                    try:
                        after=Image.open('/workspace/.agi-desktop/screen.png').copy();record.update(pixel_diff(before,after))
                        rejection=json.loads(Path('/tmp/agi-desktop/frame.json').read_text());record['frameAtRejection']=rejection['frame']
                        if self.capture and len(self.images)<6:
                            import io,base64
                            for name,im in [('before',before),('rejected',after)]:
                                data=io.BytesIO();im.save(data,format='PNG');raw=data.getvalue()
                                if len(raw)<=180000:self.images.append({'name':stage+'-'+str(number)+'-'+name+'.png','png':base64.b64encode(raw).decode()})
                    except (OSError,ValueError,KeyError,RuntimeError):
                        record['forensicsUnavailable']=True
                if not known or number==3:raise
        raise RuntimeError('SELFTEST_ACTION_ATTEMPTS_EXHAUSTED')

def memory_evidence():
 out={'childMaxRSSKiB':resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss}
 for name in ['memory.peak','memory.events','memory.max']:
  p=Path('/sys/fs/cgroup')/name
  if p.is_file():out[name]=p.read_text().strip()[:1024]
 return out

def main():
 from docx import Document
 root=Path('/workspace/selftest');root.mkdir(exist_ok=True)
 for f in [root/'typed.txt',root/'ready.json']:
  if f.exists():f.unlink()
 app=Path(__file__).with_name('fixture.py')
 state=root/'fixture-state.json'
 if state.exists():state.unlink()
 # Only this disposable synthetic-test container receives a quiet Fluxbox profile.
 wm=Path('/tmp/agi-desktop/home/.fluxbox');wm.mkdir(parents=True,exist_ok=True)
 (wm/'init').write_text('session.screen0.toolbar.visible: false\nsession.screen0.focusModel: ClickFocus\nsession.screen0.focusNewWindows: true\n')
 desktop('start');log=open(root/'fixture.log','w');p=subprocess.Popen(['python3',str(app)],env={**os.environ,'DISPLAY':':99'},stdout=log,stderr=log)
 stage='fixture-start';started=time.monotonic();peaks={};stopped=threading.Event();actions=FixtureActions('--capture-fixture' in __import__('sys').argv);durations={};stage_at=started
 def mark(next_stage):
  nonlocal stage,stage_at
  now=time.monotonic();durations[stage]=round(now-stage_at,4);stage=next_stage;stage_at=now
 def sample():
  while not stopped.wait(.1):
   totals={}
   for proc in Path('/proc').glob('[0-9]*'):
    try:
     name=(proc/'comm').read_text().strip();text=(proc/'status').read_text();rss=int(next(line.split()[1] for line in text.splitlines() if line.startswith('VmRSS:')))
     key='browser' if 'chrom' in name else 'libreoffice' if name in ['soffice.bin','oosplash'] else 'desktop-fixture' if name in ['Xvfb','fluxbox','python3'] else None
     if key:totals[key]=totals.get(key,0)+rss
    except (OSError,StopIteration,ValueError):pass
   for key,value in totals.items():peaks[key]=max(peaks.get(key,0),value)
 thread=threading.Thread(target=sample,daemon=True);thread.start()
 try:
  quiescent_frame()
  mark('focus-entry');actions.perform(stage,'click',target='entry')
  mark('keyboard');actions.perform(stage,'key',arguments=['--key','ctrl+a'],focus=True)
  mark('type-text');actions.perform(stage,'type',arguments=['--text','AGI keyboard and mouse verified'],focus=True)
  mark('save-button');actions.perform(stage,'click',target='save',focus=True,typed=True)
  mark('saved-file');end=time.monotonic()+5
  while not (root/'typed.txt').exists() and time.monotonic()<end:time.sleep(.05)
  if not (root/'typed.txt').exists() or (root/'typed.txt').read_text()!='AGI keyboard and mouse verified':raise RuntimeError('SELFTEST_TYPED_FILE_MISMATCH')
  mark('docx-render');doc=Document();doc.add_heading('AGI rendering check',0);doc.add_paragraph('Public local fixture. Total: 42.');doc.save(root/'report.docx')
  command(stage,['agi-inspect',str(root/'report.docx'),'--out',str(root/'rendered')],180)
  if 'Total: 42.' not in (root/'rendered/text.txt').read_text() or not list((root/'rendered').glob('page*.png')):raise RuntimeError('SELFTEST_RENDER_OUTPUT_MISSING')
  mark('browser');(root/'browser.html').write_text('<!doctype html><title>AGI local browser check</title><p id="result">42</p>')
  browser=command(stage,['chromium','--headless','--no-sandbox','--disable-dev-shm-usage','--disable-background-networking','--no-first-run','--user-data-dir=/tmp/agi-selftest-browser','--dump-dom',(root/'browser.html').as_uri()],60)
  if '<p id="result">42</p>' not in browser:raise RuntimeError('SELFTEST_BROWSER_OUTPUT_MISSING')
  mark('complete')
  return {'schema':'agi-workstation-check/v3','syntheticFixture':True,'attempts':actions.attempts,'fixtureImages':actions.images,'stageDurations':durations,'retryCount':sum(a['attempt']>1 for a in actions.attempts),'passed':True,'desktopKeyboardMouse':True,'screenshot':True,'officeRendered':True,'browserRendered':True,'liveInference':False,'durationSeconds':round(time.monotonic()-started,3),'sampledProcessPeakRSSKiB':peaks,'resources':memory_evidence()}
 except Exception as e:
  detail={'inner':e.detail} if isinstance(e,CheckFailure) else {'error':str(e)[:200]}
  durations[stage]=round(time.monotonic()-stage_at,4)
  print(json.dumps({'schema':'agi-workstation-check/v3','syntheticFixture':True,'passed':False,'stage':stage,**detail,'attempts':actions.attempts,'fixtureImages':actions.images,'stageDurations':durations,'durationSeconds':round(time.monotonic()-started,3),'resources':memory_evidence(),'sampledProcessPeakRSSKiB':peaks,'walletAccessed':False,'network':'none'}),flush=True)
  e.reported=True
  raise
 finally:
  stopped.set();thread.join(timeout=2);p.terminate()
  try:p.wait(timeout=5)
  except subprocess.TimeoutExpired:p.kill();p.wait()
  log.close()

if __name__=='__main__':
 try:print(json.dumps(main()))
 except Exception as e:
  if not getattr(e,'reported',False):print(json.dumps({'schema':'agi-workstation-check/v2','passed':False,**(e.detail if isinstance(e,CheckFailure) else {'stage':'initialization','error':str(e)[:200]}),'walletAccessed':False,'network':'none'}))
  raise SystemExit(1)
